"""Characterization tests pinning FactCreator behaviour before decomposition.

The snapshot test captures every fact the samples produce as an xBRL-JSON document (concept, XBRL value, entity,
period, unit, decimals, dimensions, footnotes), so any refactor that changes a value, unit, dimension or
period — not just the fact count — fails loudly. Regenerate the snapshot by
running this module directly:

    python -m tests.unitTests.xlsx_template_reader.test_fact_creator_characterization
"""

from __future__ import annotations

import json
from collections import Counter
from pathlib import Path
from typing import ClassVar, NamedTuple

import pytest
from tests.unitTests.xbrl_json_snapshot import report_to_xbrl_json

from mireport.conversionresults import ConversionResultsBuilder, Severity
from mireport.data.disclosures import VSME_DEFAULTS
from mireport.exceptions import AmbiguousComponentException
from mireport.report import InlineReport
from mireport.taxonomy import getTaxonomy, loadBuiltInTaxonomyJSON
from mireport.xbrljson_reader import XbrlJsonProcessor
from mireport.xlsx_template_reader._binder import WorkbookBinder
from mireport.xlsx_template_reader._config import ConverterConfig
from mireport.xlsx_template_reader._enumerations import (
    LabelMatch,
    eeDomainByLabel,
    getClosestEEMemberMatch,
    resolveMemberByLabel,
)
from mireport.xlsx_template_reader._fact_creator import FactCreator
from mireport.xlsx_template_reader._fact_support import (
    processNumeric,
    resolveMemberWithMessages,
)
from mireport.xlsx_template_reader._messages import Messenger
from mireport.xlsx_template_reader._reader import WorkbookReader
from mireport.xlsx_template_reader._units import UnitResolver, cleanUnitTextFromExcel
from mireport.xlsx_template_reader.processor import XlsxProcessor
from mireport.xlsx_template_reader.util import loadExcelFromPathOrFileLike

_TESTS_DIR = Path(__file__).parent.parent.parent
_REPO_ROOT = _TESTS_DIR.parent
SAMPLE_1_2_0 = _TESTS_DIR / "data" / "VSME-Digital-Template-Sample-1.2.0.xlsx"
SAMPLE_1_3_0 = (
    _REPO_ROOT / "digital-templates" / "VSME-Digital-Template-Sample-1.3.0.xlsx"
)
_SNAPSHOT_DIR = Path(__file__).parent / "snapshots"

# sample workbook -> snapshot file. 1.3.0 is the only sample with footnotes.
SNAPSHOT_CASES = {
    SAMPLE_1_2_0: _SNAPSHOT_DIR / "vsme-1.2.0-facts.json",
    SAMPLE_1_3_0: _SNAPSHOT_DIR / "vsme-1.3.0-facts.json",
}

MAX_VALUE_LENGTH = 80


def _results() -> ConversionResultsBuilder:
    return ConversionResultsBuilder(consoleOutput=False)


def _snapshotDocument(sample: Path) -> dict:
    """The sample's facts as xBRL-JSON (see tests/unitTests/xbrl_json_snapshot.py)."""
    results = _results()
    report = XlsxProcessor.from_file(sample, results, VSME_DEFAULTS).createReport()
    entryPoint = report.taxonomy.entryPoint
    return report_to_xbrl_json(report, entryPoint)


def _extras(sample: Path) -> dict:
    """What the OIM document cannot say: ix:scale per concept, and how many messages of each
    severity the conversion gave."""
    results = _results()
    report = XlsxProcessor.from_file(sample, results, VSME_DEFAULTS).createReport()
    return {
        "scales": {
            str(f.concept.qname): f.scale for f in report.facts if f.scale is not None
        },
        "messageSeverities": dict(
            sorted(Counter(m.severity.name for m in results.messages).items())
        ),
    }


# Not OIM, so not in the xBRL-JSON snapshot: the percent facts are shown x100 (ix:scale -2).
_PERCENT_SCALES = {
    "vsme:EmployeeTurnoverRate": -2,
    "vsme:PercentageGapInPayBetweenFemaleAndMaleEmployees": -2,
    "vsme:PercentageOfEmployeesCoveredByCollectiveBargainingAgreements": -2,
}
EXPECTED_EXTRAS = {
    SAMPLE_1_2_0: {
        "scales": _PERCENT_SCALES,
        "messageSeverities": {"INFO": 4, "WARNING": 14},
    },
    SAMPLE_1_3_0: {
        "scales": _PERCENT_SCALES,
        "messageSeverities": {"INFO": 5, "WARNING": 13},
    },
}


@pytest.mark.slow
class TestFactSnapshot:
    @pytest.mark.parametrize(
        "sample,snapshot",
        SNAPSHOT_CASES.items(),
        ids=[p.stem for p in SNAPSHOT_CASES],
    )
    def test_facts_match_snapshot(self, sample: Path, snapshot: Path):
        assert sample.is_file(), f"Missing sample workbook {sample}"
        assert snapshot.is_file(), (
            f"Missing snapshot {snapshot}. Generate it by running this module "
            "directly, then review and commit it."
        )
        expected = json.loads(snapshot.read_text(encoding="utf-8"))
        actual = _snapshotDocument(sample)

        assert actual["documentInfo"] == expected["documentInfo"]
        expected_facts = list(expected["facts"].items())
        actual_facts = list(actual["facts"].items())
        # Compare pairwise for a readable diff before falling back to counts.
        for exp, act in zip(expected_facts, actual_facts):
            assert act == exp
        assert len(actual_facts) == len(expected_facts)
        assert _extras(sample) == EXPECTED_EXTRAS[sample]

    @pytest.mark.parametrize(
        "snapshot",
        SNAPSHOT_CASES.values(),
        ids=[p.stem for p in SNAPSHOT_CASES.values()],
    )
    def test_the_snapshot_is_a_document_our_xbrl_json_reader_accepts(
        self, snapshot: Path
    ):
        """It is xBRL-JSON, not a private format: the reader takes every fact (footnotes aside)."""
        document = json.loads(snapshot.read_text(encoding="utf-8"))
        document["facts"] = {
            k: {name: part for name, part in fact.items() if name != "links"}
            for k, fact in document["facts"].items()
            if fact["dimensions"]["concept"] != "xbrl:note"
        }
        processor = XbrlJsonProcessor(document, _results(), strict=True)
        processor.createReport()
        assert processor.factsAdded == len(document["facts"])


# ---------------------------------------------------------------------------
# In-process FactCreator fixture (mirrors the production wiring in processor.py).
# Built on the current shipped template (1.3.0); these tests synthesize their
# own cells and discover concepts from the taxonomy, so unlike the snapshots
# they don't depend on specific workbook content.
# ---------------------------------------------------------------------------


class CreatorEnv(NamedTuple):
    creator: FactCreator
    report: InlineReport
    bindings: object
    taxonomy: object
    reader: WorkbookReader
    results: ConversionResultsBuilder


@pytest.fixture(scope="module")
def creator_env():
    wb = loadExcelFromPathOrFileLike(SAMPLE_1_3_0)
    results = _results()
    reader = WorkbookReader(wb, results)

    entry_point = reader.value(VSME_DEFAULTS["entryPoint"]).as_str()
    taxonomy = getTaxonomy(entry_point)
    report = InlineReport(taxonomy, None)
    report.addSchemaRef(entry_point)

    for period in VSME_DEFAULTS.get("periods", []):
        start = reader.value(period["start"]).as_date()
        end = reader.value(period["end"]).as_date()
        if report.addDurationPeriod(period["name"], start, end):
            report.setDefaultPeriodName(period["name"])

    bindings = WorkbookBinder(reader, taxonomy, results).bind()
    creator = FactCreator(bindings, reader, report, results, VSME_DEFAULTS)
    yield CreatorEnv(creator, report, bindings, taxonomy, reader, results)
    wb.close()


@pytest.fixture(scope="module")
def taxonomy(creator_env):
    return creator_env.taxonomy


@pytest.fixture(scope="module")
def unit_resolver(creator_env):
    config = ConverterConfig.fromDefaults(VSME_DEFAULTS, creator_env.taxonomy)
    return UnitResolver(
        creator_env.report,
        config,
        Messenger(creator_env.results),
        creator_env.reader,
        creator_env.bindings.unit_map,
    )


# ---------------------------------------------------------------------------
# Module-level helpers
# ---------------------------------------------------------------------------


class TestCleanUnitTextFromExcel:
    def test_applies_replacements(self):
        assert cleanUnitTextFromExcel("m3", {"m3": "m^3"}) == "m^3"

    def test_no_replacements_is_identity(self):
        assert cleanUnitTextFromExcel("kg", {}) == "kg"


class TestEEDomainByLabel:
    def test_rejects_non_enumeration_concept(self, taxonomy):
        not_ee = next(
            c
            for c in sorted(taxonomy.concepts, key=str)
            if not (c.isEnumerationSet or c.isEnumerationSingle)
        )
        with pytest.raises(ValueError):
            eeDomainByLabel(not_ee)

    def test_maps_member_labels_to_members(self, taxonomy):
        ee = next(
            c
            for c in sorted(taxonomy.concepts, key=str)
            if c.isEnumerationSingle and c.getEEDomain()
        )
        domain_by_label = eeDomainByLabel(ee)
        assert domain_by_label
        member = ee.getEEDomain()[0]
        label = member.getStandardLabel()
        assert domain_by_label[label][0] == member


class TestGetClosestEEMemberMatch:
    @pytest.fixture(scope="class")
    def ee_concept(self, taxonomy):
        return next(
            c
            for c in sorted(taxonomy.concepts, key=str)
            if c.isEnumerationSingle and c.getEEDomain()
        )

    def test_close_typo_matches(self, ee_concept):
        member = ee_concept.getEEDomain()[0]
        label = member.getStandardLabel()
        result = getClosestEEMemberMatch(ee_concept, label + " x")
        assert result is not None
        assert result[0] == member

    def test_garbage_returns_none(self, ee_concept):
        assert getClosestEEMemberMatch(ee_concept, "zzz qqq 12345 xyzzy") is None


class TestResolveMemberByLabel:
    """The one implementation of the exact -> configured-alias -> closest-match
    label chain that used to be pasted across FactCreator."""

    @pytest.fixture(scope="class")
    def config(self, taxonomy):
        return ConverterConfig.fromDefaults(VSME_DEFAULTS, taxonomy)

    @pytest.fixture(scope="class")
    def ee_concept(self, taxonomy):
        return next(
            c
            for c in sorted(taxonomy.concepts, key=str)
            if c.isEnumerationSingle and c.getEEDomain()
        )

    def test_exact_label(self, taxonomy, config, ee_concept):
        member = ee_concept.getEEDomain()[0]
        label = member.getStandardLabel()
        match = resolveMemberByLabel(taxonomy, config, label)
        assert match == LabelMatch(member, False, None)

    def test_configured_alias(self, taxonomy, config):
        # Find an alias whose cell value doesn't resolve directly but whose
        # target label does (some configured aliases target older taxonomy
        # labels that no longer exist).
        usable = next(
            (
                (cell_value, target)
                for cell_value, label in config.cellValuesToTaxonomyLabels.items()
                if taxonomy.resolveConcept(
                    cell_value, by_label=True, only_reportable=False
                )
                is None
                and (
                    target := taxonomy.resolveConcept(
                        label, by_label=True, only_reportable=False
                    )
                )
                is not None
            ),
            None,
        )
        if usable is None:
            pytest.skip("no configured alias resolves against this taxonomy")
        cell_value, expected = usable
        match = resolveMemberByLabel(taxonomy, config, cell_value)
        assert match is not None
        assert match.concept == expected
        assert match.viaConfiguredAlias is True

    def test_closest_match_needs_ee_concept(self, taxonomy, config, ee_concept):
        member = ee_concept.getEEDomain()[0]
        label = member.getStandardLabel()
        typo = label + " x"
        # Without the EE concept there is no fuzzy fallback.
        assert resolveMemberByLabel(taxonomy, config, typo) is None
        match = resolveMemberByLabel(taxonomy, config, typo, ee_concept=ee_concept)
        assert match is not None
        assert match.concept == member
        assert match.closestLabel is not None

    def test_no_match_returns_none(self, taxonomy, config):
        assert resolveMemberByLabel(taxonomy, config, "zzz qqq 12345") is None


class TestResolveMemberByLabelDomainScoping:
    """ee_concept / dimension scope the exact and alias lookups to the
    relevant domain (via a resolveConcept predicate), so an exact label match
    outside the domain can no longer beat the in-domain candidates.

    Two routes to a domain: an enumeration concept carries its domain
    intrinsically (getEEDomain); an explicit dimension's maximum permitted
    domain comes from taxonomy.getDomainMembersForExplicitDimension."""

    @pytest.fixture(scope="class")
    def config(self, taxonomy):
        return ConverterConfig.fromDefaults(VSME_DEFAULTS, taxonomy)

    @pytest.fixture(scope="class")
    def ee_concept(self, taxonomy):
        return next(
            c
            for c in sorted(taxonomy.concepts, key=str)
            if c.isEnumerationSingle and c.getEEDomain()
        )

    @pytest.fixture(scope="class")
    def dimension(self, taxonomy):
        dim = next(
            (
                c
                for c in sorted(taxonomy.concepts, key=str)
                if c.isExplicitDimension
                and taxonomy.getDomainMembersForExplicitDimension(c)
            ),
            None,
        )
        if dim is None:
            pytest.skip("vsme taxonomy has no explicit dimension with domain members")
        return dim

    def test_in_domain_exact_match_survives_scoping(self, taxonomy, config, ee_concept):
        member = ee_concept.getEEDomain()[0]
        match = resolveMemberByLabel(
            taxonomy, config, member.getStandardLabel(), ee_concept=ee_concept
        )
        assert match == LabelMatch(member, False, None)

    def test_out_of_domain_exact_match_is_excluded(self, taxonomy, config):
        """A label whose exact match lies outside the EE domain must not win;
        the scoped chain either finds an in-domain member or nothing."""

        def probe():
            ees = [
                c
                for c in sorted(taxonomy.concepts, key=str)
                if c.isEnumerationSingle and c.getEEDomain()
            ]
            for scope_ee in ees:
                domain = set(scope_ee.getEEDomain())
                for other in ees:
                    for foreign in other.getEEDomain():
                        if foreign in domain:
                            continue
                        label = foreign.getStandardLabel()
                        if label is None:
                            continue
                        try:
                            unscoped = resolveMemberByLabel(taxonomy, config, label)
                        except AmbiguousComponentException:
                            continue
                        if unscoped is not None and unscoped.concept == foreign:
                            return scope_ee, foreign, label
            return None

        found = probe()
        if found is None:
            pytest.skip("vsme has no out-of-domain exact-label pair to probe")
        scope_ee, foreign, label = found
        match = resolveMemberByLabel(taxonomy, config, label, ee_concept=scope_ee)
        assert match is None or (
            match.concept != foreign and match.concept in set(scope_ee.getEEDomain())
        )

    def test_dimension_scoping_resolves_domain_member(
        self, taxonomy, config, dimension
    ):
        domain = taxonomy.getDomainMembersForExplicitDimension(dimension)
        for member in sorted(domain):
            label = member.getStandardLabel()
            if label is None:
                continue
            try:
                match = resolveMemberByLabel(
                    taxonomy, config, label, dimension=dimension
                )
            except AmbiguousComponentException:
                continue
            if match is not None:
                # The label belongs to member, so a unique survivor IS member.
                assert match == LabelMatch(member, False, None)
                return
        pytest.skip("no domain-member label resolves for this dimension")

    def test_dimension_scoping_excludes_out_of_domain(
        self, taxonomy, config, dimension
    ):
        """No fuzzy fallback for dimensions: an outsider's label finds nothing."""
        domain = taxonomy.getDomainMembersForExplicitDimension(dimension)

        def probe():
            for concept in sorted(taxonomy.concepts, key=str):
                if concept in domain:
                    continue
                for label in concept.getAllStandardLabels():
                    try:
                        unscoped = resolveMemberByLabel(taxonomy, config, label)
                    except AmbiguousComponentException:
                        continue
                    if unscoped is not None and unscoped.concept == concept:
                        return label
            return None

        outsider_label = probe()
        if outsider_label is None:
            pytest.skip("no out-of-domain label resolves unscoped")
        assert (
            resolveMemberByLabel(taxonomy, config, outsider_label, dimension=dimension)
            is None
        )


class TestResolveMemberByLabelScopePlumbing:
    """Stub-level guarantees for the scoping mechanics that must hold
    regardless of the vsme taxonomy's shape."""

    class _Config:
        cellValuesToTaxonomyLabels: ClassVar[dict] = {}

    class _RecordingTaxonomy:
        def __init__(self, domain_members=frozenset()):
            self.calls = []
            self._domain_members = frozenset(domain_members)

        def getDomainMembersForExplicitDimension(self, dimension):
            return self._domain_members

        def resolveConcept(self, text, **kwargs):
            self.calls.append((text, kwargs))

    def test_ambiguity_propagates_out_of_the_chain(self):
        """The chain stays message-silent: ambiguity reaches the callers, who
        catch and report (resolveMemberWithMessages / _addExplicitDimensions)."""

        class AmbiguousTaxonomy:
            def resolveConcept(self, text, **kwargs):
                raise AmbiguousComponentException(f"'{text}' is ambiguous")

        with pytest.raises(AmbiguousComponentException):
            resolveMemberByLabel(AmbiguousTaxonomy(), self._Config(), "SharedLabel")

    def test_ee_and_dimension_are_mutually_exclusive(self):
        with pytest.raises(ValueError):
            resolveMemberByLabel(
                self._RecordingTaxonomy(),
                self._Config(),
                "Anything",
                ee_concept=object(),
                dimension=object(),
            )

    def test_ee_scope_predicate_is_passed(self):
        class _FakeMember:
            def getAllStandardLabels(self):
                return ["completely unrelated zzz"]

        in_domain = _FakeMember()
        out_of_domain = _FakeMember()

        class _FakeEE:
            isEnumerationSingle = True
            isEnumerationSet = False

            def getEEDomain(self):
                return (in_domain,)

        stub = self._RecordingTaxonomy()
        assert (
            resolveMemberByLabel(stub, self._Config(), "Anything", ee_concept=_FakeEE())
            is None
        )
        assert stub.calls
        predicate = stub.calls[0][1]["predicate"]
        assert predicate(in_domain)
        assert not predicate(out_of_domain)

    def test_dimension_scope_predicate_is_passed(self):
        in_domain = object()
        out_of_domain = object()
        stub = self._RecordingTaxonomy(domain_members={in_domain})
        assert (
            resolveMemberByLabel(stub, self._Config(), "Anything", dimension=object())
            is None
        )
        assert stub.calls
        predicate = stub.calls[0][1]["predicate"]
        assert predicate(in_domain)
        assert not predicate(out_of_domain)

    def test_empty_dimension_domain_rejects_everything(self):
        stub = self._RecordingTaxonomy(domain_members=frozenset())
        assert (
            resolveMemberByLabel(stub, self._Config(), "Anything", dimension=object())
            is None
        )
        predicate = stub.calls[0][1]["predicate"]
        assert not predicate(object())


class TestResolveMemberWithMessagesAmbiguity:
    """Ambiguity escaping the label chain must surface as a warning + None at
    the message layer, not crash the conversion."""

    def test_ambiguous_member_warns_and_returns_none(self, taxonomy, monkeypatch):
        import mireport.xlsx_template_reader._fact_support as fact_support

        class _Candidate:
            def __init__(self, qname):
                self.qname = qname

        def raiser(*args, **kwargs):
            raise AmbiguousComponentException(
                "ambiguous",
                candidates=(_Candidate("vsme:One"), _Candidate("vsme:Two")),
            )

        monkeypatch.setattr(fact_support, "resolveMemberByLabel", raiser)

        ee_concept = next(
            c
            for c in sorted(taxonomy.concepts, key=str)
            if c.isEnumerationSingle and c.getEEDomain()
        )

        class _Holder:
            def excelRef(self, cell):
                return None

        results = ConversionResultsBuilder(consoleOutput=False)
        member = resolveMemberWithMessages(
            Messenger(results),
            taxonomy,
            None,
            "Shared label",
            ee_concept,
            _Holder(),
            None,
        )
        assert member is None
        warnings = [
            str(m.messageText)
            for m in results.messages
            if m.severity is Severity.WARNING
        ]
        assert len(warnings) == 1
        assert "vsme:One" in warnings[0] and "vsme:Two" in warnings[0]
        assert "Shared label" in warnings[0]


# ---------------------------------------------------------------------------
# Unit resolution chain
# ---------------------------------------------------------------------------


def _makeCell(value, number_format=None):
    from openpyxl import Workbook

    wb = Workbook()
    cell = wb.active.cell(row=1, column=1)
    cell.value = value
    if number_format:
        cell.number_format = number_format
    return cell


class TestGetSimpleUnit:
    @pytest.fixture(scope="class")
    def any_holder(self, creator_env):
        return next(iter(creator_env.bindings.concept_map.values()))

    def test_direct_unit_id(self, unit_resolver, any_holder):
        unit = unit_resolver.getSimpleUnit(any_holder, _makeCell("MWh"))
        assert unit is not None and str(unit).endswith("MWh")

    def test_parenthesised_unit_id(self, unit_resolver, any_holder):
        unit = unit_resolver.getSimpleUnit(
            any_holder, _makeCell("Megawatt hours (MWh)")
        )
        assert unit is not None and str(unit).endswith("MWh")

    def test_unknown_unit_returns_none(self, unit_resolver, any_holder):
        assert (
            unit_resolver.getSimpleUnit(any_holder, _makeCell("wibbles per parsec"))
            is None
        )

    def test_empty_cell_returns_none(self, unit_resolver, any_holder):
        assert unit_resolver.getSimpleUnit(any_holder, _makeCell(None)) is None


class TestSetFallbackUnitForName:
    def test_non_numeric_concept_returns_false(
        self, creator_env, unit_resolver, taxonomy
    ):
        concept = next(
            c
            for c in sorted(taxonomy.concepts, key=str)
            if c.isReportable and not c.isNumeric
        )
        fb = creator_env.report.getFactBuilder().setConcept(concept)
        assert unit_resolver.setFallbackUnitForName(None, concept, fb) is False

    def test_numeric_concept_gets_a_unit(self, creator_env, unit_resolver, taxonomy):
        concept = next(
            c
            for c in sorted(taxonomy.concepts, key=str)
            if c.isReportable and c.isNumeric and not c.isMonetary
        )
        fb = creator_env.report.getFactBuilder().setConcept(concept)

        class FakeDn:
            name = "test_range"

        assert unit_resolver.setFallbackUnitForName(FakeDn(), concept, fb) is True
        assert fb._unit is not None


class TestProcessNumeric:
    def test_decimals_from_number_format(self, creator_env, taxonomy):
        report, bindings = creator_env.report, creator_env.bindings
        holder = next(iter(bindings.concept_map.values()))
        concept = next(
            c
            for c in sorted(taxonomy.concepts, key=str)
            if c.isReportable
            and c.isNumeric
            and c.dataType.localName != "percentItemType"
        )
        fb = report.getFactBuilder().setConcept(concept)
        processNumeric(
            Messenger(creator_env.results),
            holder,
            _makeCell(12.345, "0.00"),
            fb,
            12.345,
        )
        assert fb._decimals == 2

    def test_plain_format_means_inf_decimals(self, creator_env, taxonomy):
        report, bindings = creator_env.report, creator_env.bindings
        holder = next(iter(bindings.concept_map.values()))
        concept = next(
            c
            for c in sorted(taxonomy.concepts, key=str)
            if c.isReportable
            and c.isNumeric
            and c.dataType.localName != "percentItemType"
        )
        fb = report.getFactBuilder().setConcept(concept)
        processNumeric(
            Messenger(creator_env.results), holder, _makeCell(12, "General"), fb, 12
        )
        assert fb._decimals == "INF"


if __name__ == "__main__":
    loadBuiltInTaxonomyJSON()
    _SNAPSHOT_DIR.mkdir(parents=True, exist_ok=True)
    for sample, snapshot in SNAPSHOT_CASES.items():
        snapshot.write_text(
            json.dumps(_snapshotDocument(sample), indent=2, ensure_ascii=False) + "\n",
            encoding="utf-8",
        )
        print(f"Snapshot written to {snapshot}")
